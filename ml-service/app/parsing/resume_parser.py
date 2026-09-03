from pathlib import Path
from typing import Dict
import re


# =========================================================
# Optional parsers
# =========================================================

def extract_pdf_text(path: Path) -> str:
    """
    Extract text from a PDF resume.

    Requires:
        pypdf

    Install with:
        pip install pypdf
    """

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError(
            "PDF support requires pypdf. "
            "Install it with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))

    pages = []

    for page in reader.pages:

        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


def extract_docx_text(path: Path) -> str:
    """
    Extract text from a DOCX resume.

    Requires:
        python-docx

    Install with:
        pip install python-docx
    """

    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "DOCX support requires python-docx. "
            "Install it with: pip install python-docx"
        ) from exc

    document = Document(
        str(path)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Also capture text from tables because many resumes
    # put skills/contact information inside tables.
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    row_text.append(text)

            if row_text:
                paragraphs.append(
                    " | ".join(row_text)
                )

    return "\n".join(paragraphs)


# =========================================================
# Plain text support
# =========================================================

def extract_txt_text(path: Path) -> str:
    """
    Extract text from a plain-text resume.
    """

    return path.read_text(
        encoding="utf-8",
        errors="replace",
    )


# =========================================================
# Main text extraction
# =========================================================

def extract_resume_text(
    path: str | Path,
) -> str:
    """
    Extract raw text from a resume.

    Supported formats:
        .pdf
        .docx
        .txt
    """

    path = Path(path)

    if not path.exists():

        raise FileNotFoundError(
            f"Resume file not found: {path}"
        )

    if not path.is_file():

        raise ValueError(
            f"Resume path is not a file: {path}"
        )

    suffix = (
        path.suffix
        .lower()
    )

    if suffix == ".pdf":

        text = extract_pdf_text(
            path
        )

    elif suffix == ".docx":

        text = extract_docx_text(
            path
        )

    elif suffix == ".txt":

        text = extract_txt_text(
            path
        )

    else:

        raise ValueError(
            "Unsupported resume format: "
            f"{suffix}. "
            "Supported formats are PDF, DOCX, and TXT."
        )

    # -----------------------------------------------------
    # Basic cleanup
    # -----------------------------------------------------

    text = text.replace(
        "\x00",
        " ",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


# =========================================================
# Resume object
# =========================================================

def parse_resume(
    path: str | Path,
) -> Dict:
    """
    Return the first structured representation of a resume.

    IMPORTANT:

    This stage only extracts and preserves the resume text.

    Skill extraction and capability inference are handled
    separately so that the parser does not make unsupported
    claims about the user.
    """

    path = Path(path)

    raw_text = extract_resume_text(
        path
    )

    return {
        "filename": path.name,
        "file_type": path.suffix.lower(),
        "raw_text": raw_text,
    }


# =========================================================
# Standalone test
# =========================================================

if __name__ == "__main__":

    import sys

    if len(sys.argv) != 2:

        print(
            "Usage:"
        )

        print(
            "python "
            "ml-service/app/parsing/resume_parser.py "
            "<resume_path>"
        )

        raise SystemExit(1)

    resume_path = sys.argv[1]

    resume = parse_resume(
        resume_path
    )

    print()
    print("=" * 70)
    print("RESUME PARSER TEST")
    print("=" * 70)

    print()
    print("Filename:")
    print(resume["filename"])

    print()
    print("File type:")
    print(resume["file_type"])

    print()
    print("Extracted text:")
    print("-" * 70)

    print(
        resume["raw_text"]
    )

    print()
    print(
        "Characters extracted:",
        len(
            resume["raw_text"]
        ),
    )